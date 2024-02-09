### {{{                          --     import     --
from biocomp import utils as ut
import argparse
import json
import sys
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font
import xxhash

# using base58 instead of base64 because it's url-safe
import base58

import biocomp.utils as ut
import biocomp as bc

import psycopg2
from psycopg2.extras import execute_values
import logging
import os
from tqdm import tqdm

##────────────────────────────────────────────────────────────────────────────}}}

### {{{                         --     defaults     --
tlog = logging.getLogger('biocomp_tools_common')
tlog.setLevel(logging.DEBUG)

DEFAULT_CALIB_PATHS = [
    './data/calibrated_data_v3',
    './data/calibrated_data_v2',
    './data/calibrated_data',
]
DEFAULT_CALIB_NAMES = ['v3', 'v2', 'old']
DEFAULT_XP_PATH = ut.DEFAULT_XP_PATH
DEFAULT_RECIPE_PATH = ut.DEFAULT_RECIPE_PATH
DEFAULT_XP_CACHE_DIR = './devtmp/cache/xp_objs'
DEFAULT_DATA_CONFIG = {
    'network_cache_location': './__cache/network',
    'training_cache_location': './__cache/training',
    'densities_cache_location': './__cache/densities',
    'data_min_value': 500,
    'data_max_value': 100000000.0,
    'data_log_offset': 3000.0,
    'data_log_factor': 100,
    'data_log_poly_threshold': 300,
    'data_log_poly_compression': 0.4,
    'data_sampling_kde_bw_method': 0.02,
    'data_sampling_max_density_samples': 4000,
    'data_sampling_density_quantile_threshold': 0.025,
    'data_sampling_coords_for_density_threshold': 0.15,
}
DEFAULT_DATA_CONFIG_PATH = None

BIOCOMP_LOCAL_VAR_FILE = os.getenv('BIOCOMP_LOCAL_VAR_FILE', '__local_vars.py')

BIOCOMP_LOCAL_VARS = {}


def get_from_env_or_local_vars(
    varname, filename=BIOCOMP_LOCAL_VAR_FILE, default=None, raise_error=True
):
    """
    Get a variable from the environment (in priority), or from local var file if available.
    """
    global BIOCOMP_LOCAL_VARS

    if varname in os.environ:
        return os.environ[varname]

    if filename in BIOCOMP_LOCAL_VARS:
        if varname in BIOCOMP_LOCAL_VARS[filename]:
            return BIOCOMP_LOCAL_VARS[filename][varname]

    if os.path.exists(filename):
        tlog.debug(f'Loading local vars from {filename}')
        with open(filename, 'r') as f:
            local_vars = {}
            exec(f.read(), {}, local_vars)
            tlog.debug(f'Local vars found: {local_vars.keys()}')
            BIOCOMP_LOCAL_VARS[filename] = local_vars
            if varname in local_vars:
                return local_vars[varname]

    if raise_error and default is None:
        raise ValueError(f'Variable {varname} not found in either environment or {filename}')
    else:
        tlog.debug(
            f'Variable {varname} not found in either environment or {filename}, using default value'
        )
        return default


##────────────────────────────────────────────────────────────────────────────}}}


### {{{                       --     general utils     --
def parse_list(input_string):
    # Split the string by comma and then strip whitespaces from each element
    if input_string is None:
        return []
    if isinstance(input_string, list):
        return input_string
    return [element.strip() for element in input_string.split(',')]


def get_name_hash(name):
    # base58 encode the xxhash of the name
    hh = xxhash.xxh128(name).digest()
    return base58.b58encode(hh).decode('utf-8')


def filter_df(df, **filters):
    for key, value in filters.items():
        if len(df) == 0:
            return df
        if callable(value):
            df = df[df[key].apply(value)]
        else:
            df = df[df[key] == value]
    return df


##────────────────────────────────────────────────────────────────────────────}}}

### {{{                     --     general db utils     --


def connect_to_db():
    BIOCOMP_DB_NAME = get_from_env_or_local_vars('BIOCOMP_DB_NAME')
    BIOCOMP_DB_USER = get_from_env_or_local_vars('BIOCOMP_DB_USER')
    BIOCOMP_DB_PASS = get_from_env_or_local_vars('BIOCOMP_DB_PASS')
    BIOCOMP_DB_HOST = get_from_env_or_local_vars('BIOCOMP_DB_HOST')
    BIOCOMP_DB_PORT = get_from_env_or_local_vars('BIOCOMP_DB_PORT')
    try:
        conn = psycopg2.connect(
            dbname=BIOCOMP_DB_NAME,
            user=BIOCOMP_DB_USER,
            password=BIOCOMP_DB_PASS,
            host=BIOCOMP_DB_HOST,
            port=BIOCOMP_DB_PORT,
        )
    except Exception as e:
        tlog.error(f'Error connecting to database: {e}')
        raise e
    return conn


def query_to_df(query, params=None, conn=None):
    try:
        if conn is None:
            conn = connect_to_db()
        cur = conn.cursor()
        cur.execute(query, params)
        df = pd.DataFrame(cur.fetchall(), columns=[desc[0] for desc in cur.description])
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        cur.close()
        conn.close()
    return df


def table_to_df(table_name, **kwargs):
    return query_to_df(f'SELECT * FROM {table_name}', **kwargs)


def execute_query(query, params=None, conn=None):
    try:
        if conn is None:
            conn = connect_to_db()
        cur = conn.cursor()
        cur.execute(query, params)
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        cur.close()
        conn.close()


def execute_many(query, params, conn=None, dry_run=False):
    if dry_run:
        tlog.info(f'Dry run: {query} with {len(params)} values:')
        tlog.info(params)
        return
    try:
        if conn is None:
            conn = connect_to_db()
        cur = conn.cursor()
        execute_values(cur, query, params)
        conn.commit()
    except Exception as e:
        conn.rollback()
        raise e
    finally:
        cur.close()
        conn.close()


def update_table(df, table_name, key_column, conn=None, dry_run=False):
    columns = df.columns.to_list()
    query = f"""
        INSERT INTO {table_name} ({', '.join(columns)})
        VALUES %s ON CONFLICT ({key_column})
        DO UPDATE SET ({', '.join(columns)}) = ({', '.join(['EXCLUDED.' + col for col in columns])})
        """
    execute_many(query, df.values, conn, dry_run=dry_run)


##────────────────────────────────────────────────────────────────────────────}}}

### {{{                       --     collection utils     --


def get_networks_in_collections(collections):
    if not isinstance(collections, list):
        collections = [collections]
    param_placeholders = ', '.join(['%s'] * len(collections))
    query = f"SELECT * FROM collection_network WHERE collection_name IN ({param_placeholders})"
    return query_to_df(query, collections)


def create_collection(name, description):
    execute_query(
        "INSERT INTO collections (name, description) VALUES (%s, %s)", (name, description)
    )


def remove_all_from_collection(collection_name):
    execute_query("DELETE FROM collection_network WHERE collection_name = %s", (collection_name,))


def delete_collection(collection_name):
    execute_query("DELETE FROM collections WHERE name = %s", (collection_name,))


def add_networks_to_collection(collection_name, network_names):
    if not isinstance(network_names, list):
        network_names = [network_names]
    values = [(collection_name, net) for net in network_names]
    execute_many("INSERT INTO collection_network (collection_name, network_name) VALUES %s", values)


##────────────────────────────────────────────────────────────────────────────}}}


### {{{                   --     network objects loading utils     --
def get_network_row(netdf, net_name):
    if net_name not in netdf['name'].values:
        raise ValueError(f'Network id {net_name} not found in database')
    net_row = netdf[netdf['name'] == net_name]
    if len(net_row) > 1:
        raise ValueError(f'Network name {net_name} is not unique in database')
    return net_row.iloc[0]


def load_network_and_data(
    netdf, net_name, lib, path_prefix='/Users/jeandisset/Dropbox (MIT)/Biocomp'
):
    row = get_network_row(netdf, net_name)
    rfile, dfile = get_recipe_and_data_filepaths(
        row['data_file'], row['recipe_file'], path_prefix=path_prefix
    )
    networks = bc.recipe.network_from_recipe(rfile, lib, inverse='all')
    # we potentially have several networks, one for each possible inversion
    # we can use the markers to select the right one
    markers = [set(bc.recipe.escape(n.get_inverted_input_proteins())) for n in networks]
    # outputs = [set(bc.recipe.escape(networks[0].get_output_proteins())) - m for m in markers]
    target_markers = set(parse_list(row['markers']))
    escaped_target_markers = bc.recipe.escape(target_markers)
    network = networks[markers.index(escaped_target_markers)]
    X, Y = bc.recipe.get_network_XY(network, dfile, color_aliases=protein_aliases)
    return network, X, Y


def get_recipe_and_data_filepaths(data_file, recipe_file, path_prefix=''):
    # check data file present
    if pd.isna(data_file):
        raise ValueError(f'Data file information for network {net_name} is missing')
    data_file = Path(path_prefix) / data_file
    data_file = Path(data_file).resolve()
    if not Path(data_file).exists():
        raise ValueError(f'Data file {data_file} not found')
    # check recipe file present
    if pd.isna(recipe_file):
        raise ValueError(f'Recipe file information for network {net_name} is missing')
    recipe_file = Path(path_prefix) / recipe_file
    recipe_file = Path(recipe_file).resolve()
    if not Path(recipe_file).exists():
        raise ValueError(f'Recipe file {recipe_file} not found')
    return recipe_file, data_file


##────────────────────────────────────────────────────────────────────────────}}}


### {{{                        --     CLIProgram     --
class CLIProgram:
    def __init__(self):
        self.is_notebook = 'ipykernel' in sys.modules
        self.parser = argparse.ArgumentParser()

    def add_argument(self, *args, **kwargs):
        self.parser.add_argument(*args, **kwargs)

    def parse_args(self, default_args=None):
        extra_args = default_args if default_args is not None else []

        # combine parsed args and extra_args. parsed args have priority over extra_args.
        # if we're in a notebook, only use extra_args. Otherwise we can combine them.
        if self.is_notebook:
            self.args = self.parser.parse_args(extra_args)
        else:
            self.args = self.parser.parse_args(extra_args + sys.argv[1:])
            ut.logger.info(f'args: {self.args}')

        self._postprocess_args()

    def _postprocess_args(self):
        pass

    def __getattr__(self, attr):
        if attr in self.__dict__:
            return self.__dict__[attr]
        elif hasattr(self, 'args') and hasattr(self.args, attr):
            return getattr(self.args, attr)
        else:
            raise AttributeError(f"{self.__class__.__name__} object has no attribute '{attr}'")


##────────────────────────────────────────────────────────────────────────────}}}

### {{{                         --     df tools     --


def merge_update(
    left_df, right_df, key_column, priority, use_left=None, use_right=None, how='outer'
):
    """
    Merge two pandas dataframes with priority-based column selection.

    Parameters:
    priority (str): 'left' or 'right' to set priority dataframe for overlapping columns.
    use_left (list[str], optional): Columns to forcibly use from left_df.
    use_right (list[str], optional): Columns to forcibly use from right_df.

    Returns:
    pd.DataFrame: Merged dataframe based on the specified rules.
    """

    tlog.debug(
        f"""Merging dataframes with priority {priority}.
    Left columns: {left_df.columns} with shape {left_df.shape}.
    Right columns: {right_df.columns} with shape {right_df.shape}."""
    )

    use_right = use_right or []
    use_left = use_left or []

    if not set(use_left).isdisjoint(use_right):
        raise ValueError("Columns in use_left and use_right must be disjoint")

    common_columns = (
        set(left_df.columns).intersection(set(right_df.columns)).difference([key_column])
    )

    # Rename common columns in right_df to avoid suffixes in the merged dataframe
    rename_columns = {col: col + '_right' for col in common_columns if col not in use_left}
    right_df_renamed = right_df.rename(columns=rename_columns)
    merged_df = pd.merge(left_df, right_df_renamed, on=key_column, how=how)

    tlog.debug(f'Common columns: {common_columns}, renamed to {rename_columns}')

    # Apply use_left, use_right, and priority rules
    for col in common_columns:
        if col in use_right or (col + '_right' in merged_df and priority == 'right'):
            merged_df[col] = merged_df[col + '_right']

        merged_df.drop(columns=[col + '_right'], inplace=True, errors='ignore')

    return merged_df


def reorder_columns_front(df, columns):
    """
    Puts the specified columns in front of the dataframe.
    """
    columns = list(columns)
    return df[columns + [col for col in df.columns if col not in columns]]


def reorder_columns_back(df, columns):
    """
    Puts the specified columns in back of the dataframe.
    """
    columns = list(columns)
    return df[[col for col in df.columns if col not in columns] + columns]


##────────────────────────────────────────────────────────────────────────────}}}
